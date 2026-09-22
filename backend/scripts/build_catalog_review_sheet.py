"""Build the one-pass catalog review sheet for James.

Reads the seeder's product list from its AST — never from a copy — so re-running
this after the seeder changes produces a current sheet rather than a stale one.

⚠️ PRICES ARE THE POINT OF THIS FILE, AND A LITERAL-ONLY READ LOSES THEM ALL.
The seeder writes `Decimal("1405.00")`, which is an `ast.Call`. An earlier probe
used `ast.literal_eval` and reported 55 silent `None`s for the single column only
James can check. `_price_of` handles the Call form, and `_assert_prices_parsed`
fails the build if any product whose source HAS a price reads as None.

Usage:  .venv/bin/python scripts/build_catalog_review_sheet.py
Output: docs/catalog/sunnycrest-catalog-review.xlsx
"""
from __future__ import annotations

import ast
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BACKEND = pathlib.Path(__file__).resolve().parent.parent
SEEDER = BACKEND / "app" / "services" / "sunnycrest_product_seeder.py"
OUT = BACKEND.parent / "docs" / "catalog" / "sunnycrest-catalog-review.xlsx"

# ── portal mapping, from the analysis in ebd103b4 ────────────────────────
#: wilbert_sku -> portal vault id(s). Hand-authored and reviewable; never a
#: name match. The ambiguous entries are marked rather than resolved.
PORTAL_BY_SKU: dict[str, str] = {
    "BV-WBR": "wilbert-bronze", "BV-BTR": "bronze-triune",
    "BV-CTR": "copper-triune", "BV-STR": "stainless-triune",
    "BV-CR": "cameo-rose-triune", "BV-VTR": "veteran-triune",
    "BV-CON": "continental", "BV-SAL": "salute", "BV-MON": "monticello",
    "BV-MCH": "monarch", "BV-GL": "concrete-graveliner",
    "UV-BTR": "bronze-triune-urn", "UV-CTR": "copper-triune-urn",
    "UV-STR": "stainless-triune-urn", "UV-CR": "cameo-rose-triune-urn",
    "UV-VTR": "veteran-triune-urn", "UV-MON": "monticello-urn",
    "UV-SAL": "salute-urn",
}
AMBIGUOUS_BY_SKU: dict[str, str] = {
    "BV-TRB": "white-tribute AND gray-tribute — two portal vaults, one product",
    "BV-VEN": "white-venetian AND venetian — two portal vaults, one product",
    "UV-VEN": "white-venetian-urn AND venetian-urn — two portal vaults, one product",
    "UV-CG": "universal-urn — this OR White & Silver (UV-WS), not both",
    "UV-WS": "universal-urn — this OR Cream & Gold (UV-CG), not both",
}
#: Portal vaults that deliberately map to no product.
PORTAL_UNMAPPED = ("other", "cremation-other")

#: On James's list, absent from the seeder (ebd103b4 §1.2). Names verbatim.
MISSING_FROM_SEEDER: list[tuple[str, str]] = [
    ('Monticello 34"', "bought-in"), ('Large 34"', "bought-in"),
    ('Large 36"', "bought-in"), ('Large 40"', "bought-in"),
    ("Youth MT", "bought-in"), ("Youth GL", "bought-in"),
    ("Universal (P400WS)", "urn vault"), ("Basic Gray (P410)", "urn vault"),
]

NOTES_BY_NAME = {
    "Pine Box": "⚠ NO PRICE in the seeder",
    "Vault Placer": "⚠ price is 0.00",
}


def _price_of(node: ast.expr):
    """`Decimal("1405.00")` -> "1405.00"; a bare constant -> its value."""
    if isinstance(node, ast.Call) and getattr(node.func, "id", "") == "Decimal":
        if node.args and isinstance(node.args[0], ast.Constant):
            return node.args[0].value
    if isinstance(node, ast.Constant):
        return node.value
    return None


def _has_price_in_source(node: ast.expr) -> bool:
    """Did the SOURCE carry a price at all? `None` in the source is legitimate
    (Pine Box); `None` from a failed parse is the defect this separates out."""
    if isinstance(node, ast.Constant) and node.value is None:
        return False
    return True


def read_seeder() -> list[dict]:
    tree = ast.parse(SEEDER.read_text())
    rows = []
    for node in ast.walk(tree):
        if not (isinstance(node, ast.Tuple) and len(node.elts) >= 3):
            continue
        if not (isinstance(node.elts[0], ast.Constant)
                and isinstance(node.elts[0].value, str)):
            continue
        extras = {}
        if len(node.elts) > 3 and isinstance(node.elts[3], ast.Dict):
            for k, v in zip(node.elts[3].keys, node.elts[3].values):
                if isinstance(k, ast.Constant) and isinstance(v, ast.Constant):
                    extras[k.value] = v.value
        rows.append({
            "name": node.elts[0].value,
            "price": _price_of(node.elts[2]),
            "source_has_price": _has_price_in_source(node.elts[2]),
            "wilbert_sku": extras.get("wilbert_sku") or "",
            "sku": extras.get("sku") or "",
            "product_line": extras.get("product_line") or "",
        })
    return rows


def _assert_prices_parsed(rows: list[dict]) -> None:
    """⚠️ THE CHECK THE DISPATCH ASKED FOR. Fails loudly rather than shipping a
    sheet of blanks in the one column only James can fill from knowledge."""
    lost = [r["name"] for r in rows if r["source_has_price"] and r["price"] is None]
    if lost:
        raise SystemExit(
            f"PRICE PARSE FAILED for {len(lost)} product(s) whose source HAS a "
            f"price: {lost}. The sheet was not written."
        )


def _report_shape(rows: list[dict]) -> None:
    """⚠️ 55 identical values of anything is a finding about the probe first."""
    print(f"  products read           : {len(rows)}")
    for col in ("price", "wilbert_sku", "sku", "product_line"):
        distinct = len({r[col] for r in rows})
        print(f"  distinct {col:<16}: {distinct}"
              + ("   <- SUSPICIOUS, check the probe" if distinct <= 1 else ""))
    print(f"  priced in source        : {sum(r['source_has_price'] for r in rows)}")
    print(f"  no price in source      : "
          f"{sum(not r['source_has_price'] for r in rows)}")


HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="44546A")
FILL_ME = PatternFill("solid", fgColor="FFF2CC")
FILL_ME_HDR = PatternFill("solid", fgColor="BF8F00")
WARN = PatternFill("solid", fgColor="FCE4E4")
AMBIG = PatternFill("solid", fgColor="E2EFDA")

COLUMNS = [
    ("Product", 34), ("In seeder?", 11), ("wilbert_sku", 13), ("sku", 11),
    ("Product line", 17), ("Seeder price", 13), ("Portal vault id(s)", 46),
    ("Notes", 26),
    ("STILL SOLD? (Y/N)", 18), ("CURRENT PRICE", 15),
    ("MADE or BOUGHT IN", 19), ("COLOUR: separate product / one product", 40),
]
FILL_FROM = 8          # zero-based index of the first James column


def build() -> None:
    rows = read_seeder()
    _assert_prices_parsed(rows)
    _report_shape(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Catalog"

    for i, (title, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.font = HDR
        c.fill = FILL_ME_HDR if i - 1 >= FILL_FROM else HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    r = 2
    for p in rows:
        sku = p["wilbert_sku"]
        portal = PORTAL_BY_SKU.get(sku, "")
        ambiguous = AMBIGUOUS_BY_SKU.get(sku)
        ws.cell(row=r, column=1, value=p["name"])
        ws.cell(row=r, column=2, value="yes")
        ws.cell(row=r, column=3, value=sku)
        ws.cell(row=r, column=4, value=p["sku"])
        ws.cell(row=r, column=5, value=p["product_line"])
        ws.cell(row=r, column=6, value=p["price"])
        ws.cell(row=r, column=7, value=ambiguous or portal)
        ws.cell(row=r, column=8, value=NOTES_BY_NAME.get(p["name"], ""))
        if ambiguous:
            for col in (1, 7, 12):
                ws.cell(row=r, column=col).fill = AMBIG
        if p["name"] in NOTES_BY_NAME:
            ws.cell(row=r, column=8).fill = WARN
            ws.cell(row=r, column=6).fill = WARN
        for col in range(FILL_FROM + 1, len(COLUMNS) + 1):
            ws.cell(row=r, column=col).fill = FILL_ME
        r += 1

    for name, kind in MISSING_FROM_SEEDER:
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value="NOT IN SEEDER")
        ws.cell(row=r, column=5, value=kind)
        ws.cell(row=r, column=8, value="on James's list, absent from the seeder")
        for col in (1, 2, 8):
            ws.cell(row=r, column=col).fill = WARN
        for col in range(FILL_FROM + 1, len(COLUMNS) + 1):
            ws.cell(row=r, column=col).fill = FILL_ME
        r += 1

    _readme(wb, len(rows), r - 2)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"  wrote {OUT.relative_to(BACKEND.parent)}  ({r - 2} rows)")


def _readme(wb: Workbook, seeded: int, total: int) -> None:
    ws = wb.create_sheet("README", 0)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 104
    lines = [
        ("Sunnycrest catalog review", ""),
        ("", ""),
        ("What this is",
         f"Every product the seeder would create ({seeded}), plus products on your list the "
         f"seeder does not have ({total - seeded}). {total} rows."),
        ("Why",
         "The seeder has never run on production. Before it does, the catalog needs one pass "
         "from you. Once filled, this file becomes the seeder's source, so the catalog has "
         "ONE definition you can edit."),
        ("", ""),
        ("FILL THESE IN", "the four amber columns. Nothing in them is pre-filled — no value "
                          "here is invented."),
        ("STILL SOLD? (Y/N)", "N removes it. The seeder will not create it."),
        ("CURRENT PRICE", "Blank means the seeder price is correct. A value replaces it."),
        ("MADE or BOUGHT IN",
         "The seeder never recorded this — there is no source field on the product. "
         "Write 'made' or 'bought'."),
        ("COLOUR column",
         "Only on the green rows. The portal treats some colours as separate vaults and the "
         "seeder joins them; for Universal it is the reverse. Answer 'separate product' or "
         "'one product, colour is a choice'. The practical test: different price or made "
         "differently means separate products."),
        ("", ""),
        ("Green rows", "Ambiguous: the portal vault cannot be mapped to one product until the "
                       "colour question is answered. Availability cannot be imported for these."),
        ("Pink rows", "Need attention: no price, a zero price, or absent from the seeder."),
        ("", ""),
        ("Portal vault id(s)",
         "Which vault in the ordering portal this product is, used to import personalization "
         "availability. Blank means no portal vault — equipment, fees, urns."),
        ("Portal 'Other'",
         "The portal's two 'Other' entries map to no product, deliberately."),
        ("", ""),
        ("Not answered here",
         "Whether the seeder's bought-in sizes (Continental 34\", Graveliner 34\"/38\") should "
         "exist at all — they are not on your list. Mark them N if they should not."),
    ]
    for i, (a, b) in enumerate(lines, start=1):
        ca = ws.cell(row=i, column=1, value=a)
        cb = ws.cell(row=i, column=2, value=b)
        ca.font = Font(bold=True)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            ca.font = Font(bold=True, size=14)
    ws.cell(row=6, column=1).fill = FILL_ME


if __name__ == "__main__":
    build()
